"""Corpus dorado de documentos y adjuntos: PDF, Word, @ruta e imágenes.

Fabrica los archivos con PyMuPDF/python-docx/Pillow, ejecuta el código Python
(`documents.pdf_text`, `documents.docx_text`, `commands.parse_attachments`,
`attachments_block`, `encode_image`, `parse_image_markers`) y guarda la salida.
Las implementaciones portadas (Go) reconstruyen los mismos archivos desde
fixtures/documents_corpus.json y comparan.

El texto de un PDF depende de la biblioteca de extracción: los casos marcados
"words" se comparan palabra a palabra, ignorando espacios y saltos de línea.

    pip install pypdf python-docx pymupdf pillow
    python apps/cli/validation/gen_documents_corpus.py          # reescribe
    python apps/cli/validation/gen_documents_corpus.py --check  # falla si cambió
"""
import base64
import io
import json
import struct
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import docx  # noqa: E402
import fitz  # noqa: E402
from PIL import Image  # noqa: E402

from lixbon_cli import clipboard, commands, documents  # noqa: E402

CORPUS = HERE / "fixtures" / "documents_corpus.json"
ARIAL = Path("C:/Windows/Fonts/arial.ttf")
EDGE = Path("C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe")

REPORT_HTML = """<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Informe</title>
<style>body{font-family:Arial;font-size:12pt} table{border-collapse:collapse} td,th{border:1px solid #000;padding:4px}</style></head><body>
<h1>Informe de resultados</h1>
<p>Este es el primer párrafo del informe, con acentos: canción, señal, ¿qué tal? ¡Perfecto! y una línea larga que se ajusta automáticamente al ancho de la página para probar el salto de línea natural del navegador.</p>
<p>Segundo párrafo con <b>negrita</b>, <i>cursiva</i> y un enlace <a href="https://ejemplo.com">ejemplo.com</a>.</p>
<ul><li>Elemento uno</li><li>Elemento dos</li><li>Elemento tres</li></ul>
<table><tr><th>Nombre</th><th>Cantidad</th><th>Precio</th></tr><tr><td>Manzana</td><td>3</td><td>1.20</td></tr><tr><td>Pera</td><td>10</td><td>0.80</td></tr></table>
<div style="page-break-before:always"></div><h1>Segunda página</h1><p>Texto en la segunda página.</p>
</body></html>"""
REPORT_CONTAINS = [
    "Informe de resultados",
    "Este es el primer párrafo del informe, con acentos: canción, señal, ¿qué tal? ¡Perfecto!",
    "larga que se ajusta automáticamente al ancho de la página para probar el salto de línea natural del navegador.",
    "Segundo párrafo con negrita, cursiva y un enlace ejemplo.com.",
    "Elemento uno", "Elemento dos", "Elemento tres",
    "Manzana", "1.20", "Pera", "0.80",
    "── página 2 ──", "Segunda página", "Texto en la segunda página.",
]


def pdf_from_browser() -> bytes | None:
    """Un PDF como los de verdad (fuentes TrueType subconjunto con ToUnicode)."""
    if not EDGE.exists():
        return None
    with tempfile.TemporaryDirectory() as tmp:
        html, out = Path(tmp) / "informe.html", Path(tmp) / "informe.pdf"
        html.write_text(REPORT_HTML, encoding="utf-8")
        subprocess.run([str(EDGE), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={out}", html.as_uri()], check=True, capture_output=True, timeout=120)
        return out.read_bytes()


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def packed(data: bytes) -> str | dict:
    """Los archivos de relleno largos se guardan como {fill, size}."""
    if len(data) > 4096 and len(set(data)) == 1:
        return {"fill": chr(data[0]), "size": len(data)}
    return b64(data)


def png(color: tuple[int, int, int], size: tuple[int, int] = (4, 4)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", size, color).save(out, "PNG")
    return out.getvalue()


def pdf_pages(pages: list[str], fontname: str = "helv") -> bytes:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontname=fontname, fontsize=12)
    return doc.tobytes(deflate=True, garbage=3)


def pdf_embedded_font() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_font(fontname="F0", fontfile=str(ARIAL))
    page.insert_text((72, 72), "Fuente incrustada: señal, canción y ¿qué?", fontname="F0", fontsize=12)
    doc.subset_fonts()
    return doc.tobytes(deflate=True, garbage=3)


def pdf_scan() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(72, 72, 272, 172), stream=png((200, 30, 30), (50, 25)))
    return doc.tobytes()


def pdf_table() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    for row, cells in enumerate([("Nombre", "Cantidad", "Precio"), ("Manzana", "3", "1.20"), ("Pera", "10", "0.80")]):
        for col, cell in enumerate(cells):
            page.insert_text((72 + col * 150, 100 + row * 20), cell, fontname="helv", fontsize=11)
    return doc.tobytes(deflate=True)


def pdf_encrypted() -> bytes:
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "secreto", fontname="helv", fontsize=12)
    return doc.tobytes(encryption=fitz.PDF_ENCRYPT_RC4_128, owner_pw="dueno", user_pw="usuario")


def docx_bytes(build) -> bytes:
    document = docx.Document()
    build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def docx_template() -> bytes:
    """El .docx vacío de python-docx: los casos solo cambian word/document.xml."""
    return docx_bytes(lambda d: None)


def document_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read("word/document.xml").decode("utf-8")


def docx_paragraphs(d) -> None:
    d.add_paragraph("Primer párrafo con acentos: señal, canción.")
    d.add_paragraph("   ")
    d.add_paragraph("Segundo párrafo")


def docx_tables(d) -> None:
    d.add_paragraph("Antes de la tabla")
    table = d.add_table(rows=3, cols=3)
    for r, row in enumerate([("Nombre", "Cantidad", ""), ("Manzana", "3", "1.20"), ("", "", "")]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    d.add_paragraph("Después de la tabla")
    second = d.add_table(rows=1, cols=2)
    second.cell(0, 0).text = "A"
    second.cell(0, 1).text = "B\nlínea dos"


def docx_merged(d) -> None:
    table = d.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Unidas"
    table.cell(0, 2).text = "Sola"
    table.cell(1, 0).merge(table.cell(1, 2)).text = "Toda la fila"


def docx_vmerged(d) -> None:
    table = d.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Vertical"
    table.cell(0, 1).text = "x"
    table.cell(1, 1).text = "y"
    table.cell(2, 0).text = "z"
    table.cell(2, 1).text = "w"


def docx_cell_paragraphs(d) -> None:
    table = d.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.text = "primero"
    cell.add_paragraph("segundo")
    table.cell(0, 1).text = "otro"


def docx_runs(d) -> None:
    p = d.add_paragraph()
    p.add_run("Negrita ").bold = True
    p.add_run("y cursiva").italic = True
    p.add_run("\ttab y").add_break()
    p.add_run("salto")


def docx_empty(d) -> None:
    d.add_paragraph("")


def run_pdf(case: dict, data: bytes) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "doc.pdf"
        path.write_bytes(data)
        try:
            case["expected"] = documents.pdf_text(path)
        except Exception as exc:  # el texto del error es de pypdf: solo importa que falle
            case["error"] = f"{type(exc).__name__}: {exc}"


def run_docx(case: dict, data: bytes) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "doc.docx"
        path.write_bytes(data)
        try:
            case["expected"] = documents.docx_text(path)
        except Exception as exc:
            case["error"] = f"{type(exc).__name__}: {exc}"


def build_extraction() -> list[dict]:
    pdfs = [
        ("pdf_simple", pdf_pages(["Hola mundo\nSegunda línea"]), "exact"),
        ("pdf_español", pdf_pages(["Canción señal ¿qué? ¡hola! áéíóú ÁÉÍÓÚ ñ Ñ ü"]), "exact"),
        ("pdf_tres_paginas", pdf_pages(["Uno", "Dos", "Tres\ny más"]), "exact"),
        ("pdf_pagina_vacia_en_medio", pdf_pages(["Antes", "", "Después"]), "exact"),
        ("pdf_tabla", pdf_table(), "words"),
        ("pdf_escaneo", pdf_scan(), "exact"),
        ("pdf_205_paginas", pdf_pages([f"Página número {i}" for i in range(1, 206)]), "exact"),
        ("pdf_cifrado", pdf_encrypted(), "exact"),
        ("pdf_corrupto", b"%PDF-1.4\nesto no es un pdf\n", "exact"),
        ("pdf_no_es_pdf", b"texto plano con extension pdf", "exact"),
    ]
    if ARIAL.exists():
        pdfs.append(("pdf_fuente_incrustada", pdf_embedded_font(), "words"))
    docxs = [
        ("docx_parrafos", docx_bytes(docx_paragraphs)),
        ("docx_tablas", docx_bytes(docx_tables)),
        ("docx_celdas_combinadas", docx_bytes(docx_merged)),
        ("docx_combinacion_vertical", docx_bytes(docx_vmerged)),
        ("docx_celda_varios_parrafos", docx_bytes(docx_cell_paragraphs)),
        ("docx_runs_tab_salto", docx_bytes(docx_runs)),
        ("docx_vacio", docx_bytes(docx_empty)),
        ("docx_corrupto", b"PK\x03\x04esto no es un docx"),
        ("docx_no_es_zip", b"texto plano"),
    ]
    cases = []
    report = pdf_from_browser()
    if report:
        cases.append({"name": "pdf_navegador", "kind": "pdf", "file": b64(report), "compare": "contains",
                      "contains": REPORT_CONTAINS, "expected": None})
    for name, data, compare in pdfs:
        case = {"name": name, "kind": "pdf", "file": b64(data), "compare": compare}
        run_pdf(case, data)
        if name == "pdf_fuente_incrustada":
            case["known_gap"] = ("la biblioteca Go no lee el ToUnicode con líneas rotas que genera PyMuPDF; "
                                 "pypdf sí. Comprobado que Edge/Word no lo producen.")
        cases.append(case)
    for name, data in docxs:
        case = {"name": name, "kind": "docx", "compare": "exact"}
        if data.startswith(b"PKesto") or not data.startswith(b"PK"):
            case["file"] = b64(data)
        else:
            case["document_xml"] = document_xml(data)
        run_docx(case, data)
        cases.append(case)
    return cases


ATTACH_FILES: dict[str, bytes] = {
    "notas.txt": "línea uno\r\nlínea dos\r\n".encode("utf-8"),
    "con espacios.txt": "contenido con espacios\n".encode("utf-8"),
    "código.py": "print('hola')\n".encode("utf-8"),
    "sub/dentro.md": "# Título\n".encode("utf-8"),
    "fence.md": "antes\n```py\nx = 1\n```\ndespués\n".encode("utf-8"),
    "binario.bin": b"\xff\xfe\x00\x01\x02",
    "latin1.txt": "canci\xf3n".encode("latin-1"),
    "grande.txt": b"a" * (64 * 1024 + 1),
    "justo.txt": b"b" * (64 * 1024),
    "logo.png": png((10, 120, 200)),
    "foto.JPG": png((255, 0, 0)),
    "doc.pdf": pdf_pages(["Contenido del PDF"]),
    "vacio.txt": b"",
}

ATTACH_CASES: list[tuple[str, str]] = [
    ("sin_adjuntos", "hola mundo"),
    ("texto_simple", "mira @notas.txt por favor"),
    ("crlf_normalizado", "@notas.txt"),
    ("ruta_con_espacios_entre_comillas", 'lee @"con espacios.txt" ahora'),
    ("unicode", "revisa @código.py"),
    ("subdirectorio", "@sub/dentro.md"),
    ("puntuacion_pegada", "mira @notas.txt, y @logo.png."),
    ("handle_de_usuario", "hola @usuario cómo estás"),
    ("email_no_es_adjunto", "escribe a yo@ejemplo.com"),
    ("directorio", "@sub"),
    ("imagen", "describe @logo.png"),
    ("imagen_mayusculas", "@foto.JPG"),
    ("imagen_inexistente", "mira @fantasma.png"),
    ("imagen_inexistente_sola", "@fantasma.png"),
    ("solo_adjunto", "@notas.txt"),
    ("dos_archivos_y_imagen", "@notas.txt @código.py y @logo.png"),
    ("misma_imagen_dos_veces", "@logo.png @logo.png"),
    ("binario", "@binario.bin"),
    ("no_utf8", "@latin1.txt"),
    ("supera_64_kB", "@grande.txt"),
    ("justo_64_kB", "@justo.txt"),
    ("pdf", "resume @doc.pdf"),
    ("con_fence", "@fence.md"),
    ("archivo_vacio", "@vacio.txt"),
    ("arroba_suelta", "a @ b"),
    ("comillas_sin_cierre", '@"notas.txt'),
    ("parentesis", "(mira @notas.txt)"),
]


def build_attachments() -> list[dict]:
    cases = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp).resolve()
        for rel, data in ATTACH_FILES.items():
            target = root / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(data)
        for name, text in ATTACH_CASES:
            clean, images, files, errors = commands.parse_attachments(text, root)
            cases.append({
                "name": name,
                "text": text,
                "expected": {
                    "clean": clean,
                    "images": [p.relative_to(root).as_posix() for p in images],
                    "files": [[n, c] for n, c in files],
                    "errors": errors,
                },
            })
    return cases


def build_blocks() -> list[dict]:
    samples = [
        [["a.txt", "uno\ndos\n"]],
        [["a.txt", "uno"], ["b.txt", "dos\n\n"]],
        [["fence.md", "antes\n```py\nx\n```\n"]],
        [],
    ]
    return [{"files": s, "expected": commands.attachments_block([tuple(f) for f in s])} for s in samples]


def build_encode() -> list[dict]:
    cases = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        for rel, data in [("a.png", png((1, 2, 3))), ("b.WEBP", b"RIFFxxxxWEBP"), ("c.gif", b"GIF89a"), ("d.txt", b"x")]:
            (root / rel).write_bytes(data)
        for rel in ["a.png", "b.WEBP", "c.gif", "d.txt", "no_existe.png"]:
            case = {"name": rel, "files": {k: b64((root / k).read_bytes()) for k in ["a.png", "b.WEBP", "c.gif", "d.txt"]}, "path": rel}
            try:
                case["expected"] = commands.encode_image(root / rel)
            except ValueError as exc:
                case["error"] = str(exc).replace(str(root), "{ROOT}")
            cases.append(case)
    return cases


def build_markers() -> list[dict]:
    staged = [Path("uno.png"), Path("dos.png"), Path("tres.png")]
    texts = ["[IMG#1] mira", "[IMG#3] y [IMG#1]", "[img#2]", "[IMG#9]", "[IMG#1] [IMG#1]", "sin marcador", "[IMG#0]"]
    return [{"text": t, "expected": [p.name for p in commands.parse_image_markers(t, staged)]} for t in texts]


def make_dib(width: int, height: int, bits: int, pixels: list[tuple[int, int, int]], *,
             compression: int = 0, header: int = 40) -> bytes:
    """DIB de prueba. pixels va en el orden de las filas del archivo (BGR a mano)."""
    step = bits // 8
    stride = ((width * bits + 31) // 32) * 4
    body = b""
    for y in range(abs(height)):
        row = b""
        for x in range(width):
            r, g, b = pixels[y * width + x]
            row += bytes((b, g, r)) + (b"\xff" if step == 4 else b"")
        body += row + b"\x00" * (stride - len(row))
    head = struct.pack("<IiiHHIIiiII", header, width, height, 1, bits, compression, len(body), 0, 0, 0, 0)
    head += b"\x00" * (header - 40)
    if compression == 3 and header == 40:
        head += struct.pack("<III", 0xFF0000, 0x00FF00, 0x0000FF)
    return head + body


def build_dib() -> list[dict]:
    px = [(255, 0, 0), (0, 255, 0), (0, 0, 255), (10, 20, 30), (200, 100, 50), (1, 2, 3)]
    samples = [
        ("24_bits_de_abajo_a_arriba", make_dib(3, 2, 24, px)),
        ("24_bits_de_arriba_a_abajo", make_dib(3, -2, 24, px)),
        ("32_bits_alfa_descartado", make_dib(3, 2, 32, px)),
        ("32_bits_bitfields", make_dib(3, 2, 32, px, compression=3)),
        ("cabecera_v5", make_dib(3, 2, 24, px, header=124)),
        ("una_columna", make_dib(1, 2, 24, px[:2])),
        ("8_bits_no_soportado", make_dib(3, 2, 24, px)[:14] + struct.pack("<H", 8) + make_dib(3, 2, 24, px)[16:]),
        ("comprimido_no_soportado", make_dib(3, 2, 24, px)[:16] + struct.pack("<I", 1) + make_dib(3, 2, 24, px)[20:]),
        ("truncado", make_dib(3, 2, 24, px)[:-5]),
        ("cabecera_corta", b"\x00" * 20),
    ]
    cases = []
    for name, dib in samples:
        png = clipboard.dib_to_png(dib)
        case = {"name": name, "dib": b64(dib)}
        if png is None:
            case["png"] = None
        else:
            image = Image.open(io.BytesIO(png)).convert("RGB")
            case["width"], case["height"] = image.size
            case["rgb"] = b64(image.tobytes())
        cases.append(case)
    return cases


def build() -> dict:
    return {
        "max_pdf_pages": documents.MAX_PDF_PAGES,
        "max_image_bytes": commands.MAX_IMAGE_BYTES,
        "max_text_attachment_bytes": commands.MAX_TEXT_ATTACHMENT_BYTES,
        "files": {rel: packed(data) for rel, data in ATTACH_FILES.items()},
        "docx_template": b64(docx_template()),
        "extraction": build_extraction(),
        "attachments": build_attachments(),
        "blocks": build_blocks(),
        "encode_image": build_encode(),
        "markers": build_markers(),
        "dib": build_dib(),
    }


def main() -> int:
    text = json.dumps(build(), ensure_ascii=False, indent=1) + "\n"
    if "--check" in sys.argv:
        current = CORPUS.read_text(encoding="utf-8") if CORPUS.exists() else ""
        # Los PDF/docx incrustan fechas y ids: solo se comparan las salidas esperadas.
        old, new = json.loads(current or "{}"), json.loads(text)
        for section in ("attachments", "blocks", "markers", "dib"):
            if old.get(section) != new[section]:
                print(f"{section}: cambió", file=sys.stderr)
                return 1
        old_x = [(c["name"], c.get("expected"), "error" in c) for c in old.get("extraction", [])]
        new_x = [(c["name"], c.get("expected"), "error" in c) for c in new["extraction"]]
        if old_x != new_x:
            print("extraction: cambió", file=sys.stderr)
            return 1
        return 0
    CORPUS.write_text(text, encoding="utf-8")
    print(f"escrito {CORPUS} ({len(text) // 1024} kB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
